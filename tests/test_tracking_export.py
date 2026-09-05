# -*- coding: utf-8 -*-
"""`/tracking/export.xlsx|pdf` — ekrandaki listenin dosya hali.

Bu ucun tek isi vardir ve kolayca sessizce bozulur: DOSYA, EKRANDAKI SUZGECLE
AYNI SATIRLARI icermeli. Suzgec okunmazsa kullanici 3 satirlik bir liste suzup
986 satirlik bir Excel indirir ve bunu ancak birine yolladiktan sonra fark eder.
"""
import io
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tracking.db import ShipmentsDB
from web import deps, exports
from web.routers import tracking as tracking_router


@pytest.fixture
def db(tmp_path, monkeypatch):
    d = ShipmentsDB(path=tmp_path / "shipments.db")
    d.add_parcel(tracking_no="38120177000001", channel="NL", status="delivered",
                 consignee_name="AHA Austria", country="AT", season="FA26")
    d.add_parcel(tracking_no="38120177000002", channel="NL", status="in_transit",
                 consignee_name="Bella Germany", country="DE", season="FA26")
    monkeypatch.setattr(deps, "_shipments_db", d)
    return d


@pytest.fixture
def client(db):
    app = FastAPI()
    app.include_router(tracking_router.router)
    return TestClient(app)


def _sheet(response):
    return load_workbook(io.BytesIO(response.content)).active


def test_the_list_comes_down_as_a_spreadsheet(client):
    r = client.get("/tracking/export.xlsx")

    assert r.status_code == 200
    assert r.content[:2] == b"PK"  # xlsx bir zip'tir
    assert "takip-listesi-" in r.headers["content-disposition"]


def test_the_list_comes_down_as_a_pdf(client):
    r = client.get("/tracking/export.pdf")

    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"


def test_the_file_carries_only_the_filtered_rows(client):
    """Dosyanin tek sozu bu: ekranda ne suzuldiyse o insin."""
    sheet = _sheet(client.get("/tracking/export.xlsx?status=delivered"))
    numbers = {row[1].value for row in sheet.iter_rows(min_row=2)}

    assert numbers == {"38120177000001"}


def test_the_search_box_is_honoured_too(client):
    sheet = _sheet(client.get("/tracking/export.xlsx?search=Bella"))
    numbers = {row[1].value for row in sheet.iter_rows(min_row=2)}

    assert numbers == {"38120177000002"}


def test_the_column_filters_are_honoured_too(client):
    """Sutun basligindaki huni menusu `f_<alan>` yollar — sol panelle ayni ad."""
    sheet = _sheet(client.get("/tracking/export.xlsx?f_country=AT"))
    numbers = {row[1].value for row in sheet.iter_rows(min_row=2)}

    assert numbers == {"38120177000001"}


def test_the_status_is_written_the_way_the_screen_says_it(client):
    """Dosyayi acan kisi `in_transit` degil, panelde gordugu metni bekler."""
    sheet = _sheet(client.get("/tracking/export.xlsx?status=in_transit"))
    row = next(sheet.iter_rows(min_row=2, values_only=True))

    assert "in_transit" not in row


def test_country_and_zip_get_their_own_columns(client):
    """Ekranda alici adinin altina sikismislar; Excel'de suzulebilmeliler."""
    sheet = _sheet(client.get("/tracking/export.xlsx"))
    headers = [c.value for c in sheet[1]]

    assert len(set(headers)) == len(headers)  # tekrar eden baslik yok
    assert sheet.max_column == len(tracking_router.EXPORT_COLUMNS)


# Sayfa numarasi BILEREK disarida: ekranda 100 satir var, dosyada suzgece uyan
# HER satir olmali (`paged=False`). Diger her alan iki listede de bulunmali.
NOT_EXPORTED = {"[name='page']"}


def test_the_export_button_asks_for_the_same_fields_as_the_table():
    """30 saniyelik tuzagin kardesi: tablo bir suzgeci yolluyor da dosya
    yollamiyorsa, kullanici suzup indirdigi seyi bulamaz."""
    page = Path("web/templates/tracking.html").read_text(encoding="utf-8")
    include = page.split('hx-include="')[1].split('"')[0]

    for field in include.split(","):
        if field.strip() in NOT_EXPORTED:
            continue
        assert field.strip() in page.split("EXPORT_FIELDS")[1]


# --------------------------------------------------------- cok sayfali yazici

def test_a_workbook_can_carry_several_sheets():
    """Musteri bazli koli icerigi uc sayfa istiyor; ayri bir yazici yazilmadi."""
    book = load_workbook(io.BytesIO(exports.to_xlsx_multi([
        ("Ozet", ["A"], [[1]]),
        ("Icerik", ["B"], [[2]]),
        ("Matris", ["C"], [[3]]),
    ])))

    assert [s.title for s in book.worksheets] == ["Ozet", "Icerik", "Matris"]


def test_a_sheet_name_excel_refuses_is_repaired():
    """Excel sayfa adi 31 karakteri asamaz ve `[]:*?/\\` kabul etmez."""
    book = load_workbook(io.BytesIO(exports.to_xlsx_multi([
        ("A/B:C" + "x" * 40, ["A"], [[1]])])))
    title = book.worksheets[0].title

    assert len(title) <= 31
    assert not set(title) & set("[]:*?/\\")


def test_the_single_sheet_writer_still_behaves_the_same():
    """`to_xlsx` artik `to_xlsx_multi`ye devrediyor; basligi ve dondurulmus
    ust satiri kaybetmemeli."""
    sheet = load_workbook(io.BytesIO(
        exports.to_xlsx(["Takip No"], [["38120177000001"]], "Liste"))).active

    assert sheet.title == "Liste"
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].font.bold


def test_quantities_stay_numbers_so_excel_can_add_them_up():
    """Her hucre `str()`ten gecerken adet sutunu metin oluyordu ve alta toplam
    alinamiyordu — kullanicinin "exelde de sikinti cikariyor" dedigi buydu."""
    sheet = load_workbook(io.BytesIO(
        exports.to_xlsx(["Adet"], [[5]], "Liste"))).active

    assert sheet["A2"].value == 5


# ------------------------------------------------------------- secim sutunu

def test_the_tracking_list_has_a_tick_box_but_the_quick_search_does_not(client):
    """Panelin dar hizli arama tablosundan toplu indirme yapilmiyor."""
    assert 'class="row-pick"' in client.get("/tracking/rows").text
    assert 'class="row-pick"' not in client.get("/tracking/quick?q=AHA").text


def test_the_selection_survives_the_thirty_second_refresh():
    """Tablo 30 saniyede bir satirlari degistiriyor ve DOM'daki tikleri siliyor.
    Dogru kaynak JS kumesi olmali, tikler swap sonrasi geri konmali."""
    page = Path("web/templates/tracking.html").read_text(encoding="utf-8")

    assert "htmx:afterSwap" in page
    assert "syncPickBoxes" in page


def test_the_screen_and_the_endpoint_share_one_limit(client):
    """Iki sayi ayri yerde tutulursa biri degisip digeri kalir.

    Sablonda ARANMAZ, CIZILMIS sayfada aranir: Jinja tanimsiz bir degiskeni
    sessizce bos basar ve `const PICK_LIMIT = ;` sayfanin TUM betigini
    goturur — hicbir dugme calismaz, konsolda tek satir hata cikmaz.
    """
    page = client.get("/tracking").text

    assert f"PICK_LIMIT = {tracking_router.CONTENTS_BULK_LIMIT};" in page
