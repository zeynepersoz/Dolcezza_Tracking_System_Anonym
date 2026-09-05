# -*- coding: utf-8 -*-
"""Gunluk ozet — veri, sablon, ek, tasiyici, zamanlayici.

HICBIR test aga cikmaz: `requests.post` monkeypatch'lenir ya da `IS_MOCK`
korumasi devrededir. Bir testin gercekten mail atmasi felakettir.
"""
import asyncio
import io
import zipfile
from datetime import datetime, timedelta

import pytest

import i18n
from gls_api import config
from notify import digest, excel, mail
from notify.scheduler import JOB_ID, DailyDigest
from tracking.db import ShipmentsDB, delivered_note


@pytest.fixture(autouse=True)
def _clean_config():
    yield
    config.clear_overrides()


@pytest.fixture
def db(tmp_path):
    return ShipmentsDB(path=tmp_path / "notify.db")


def _now_iso(hours_ago: float = 0) -> str:
    return (datetime.now() - timedelta(hours=hours_ago)).isoformat(timespec="seconds")


def _event(db, tracking_no: str, status: str, at: str) -> None:
    """Olayi dogrudan yazar — `update_status` zamani kendi koyar, biz secmeliyiz."""
    pid = db.conn.execute("SELECT id FROM parcels WHERE tracking_no = ?",
                          (tracking_no,)).fetchone()["id"]
    db.conn.execute("INSERT INTO status_events (parcel_id, status, note, at) "
                    "VALUES (?, ?, '', ?)", (pid, status, at))
    db.conn.commit()


def _parcel(db, status: str = "created", **kw) -> None:
    """Parca ekler ve `add_parcel`'in otomatik yazdigi 'created' olayini siler —
    testler olay zamanlarini kendileri secmeli."""
    db.add_parcel(tracking_no=kw.pop("tracking_no", "38120177000001"),
                  channel=kw.pop("channel", "NL"), status=status, **kw)
    db.conn.execute("DELETE FROM status_events")
    db.conn.commit()


# ------------------------------------------------------------- delivered_since
def test_delivered_since_skips_older_events(db):
    _parcel(db, status="delivered")
    _event(db, "38120177000001", "delivered", _now_iso(hours_ago=72))
    assert db.delivered_since(_now_iso(hours_ago=24)) == []


def test_delivered_since_finds_recent(db):
    _parcel(db, status="delivered")
    _event(db, "38120177000001", "delivered", _now_iso(hours_ago=2))
    rows = db.delivered_since(_now_iso(hours_ago=24))
    assert [r["tracking_no"] for r in rows] == ["38120177000001"]


def test_delivered_twice_appears_once(db):
    """Ayni parca icin birden fazla teslimat olayi yazilmis olabilir."""
    _parcel(db, status="delivered")
    _event(db, "38120177000001", "delivered", _now_iso(hours_ago=3))
    _event(db, "38120177000001", "delivered", _now_iso(hours_ago=1))
    assert len(db.delivered_since(_now_iso(hours_ago=24))) == 1


def test_delivered_since_ignores_other_statuses(db):
    _parcel(db, status="in_transit")
    _event(db, "38120177000001", "in_transit", _now_iso(hours_ago=1))
    assert db.delivered_since(_now_iso(hours_ago=24)) == []


# ------------------------------------------------------------------- gather
def test_gather_survives_empty_db(db):
    data = digest.gather(db)
    assert data["total"] == 0
    assert data["delivered"] == [] and data["problem_total"] == 0
    assert data["detail_rows"] == 0


def test_gather_splits_problem_kinds(db):
    _parcel(db, status="exception")
    data = digest.gather(db, stale_days=7)
    assert len(data["exceptions"]) == 1
    assert data["stale"] == []


def test_gather_is_cumulative_not_last_24h(db):
    """Ek ve basliktaki sayilar sezonun TAMAMI olmali."""
    _parcel(db, tracking_no="38120177000001", status="delivered")
    _parcel(db, tracking_no="38120177000002", status="in_transit")
    data = digest.gather(db)
    assert data["total"] == 2
    assert data["detail_rows"] == 2 and len(data["parcels"]) == 2
    assert data["delivered_pct"] == 50
    assert data["delivered"] == []          # son 24 saatte olay yok


def test_gather_reads_season_from_view_name(db):
    config.apply_overrides({"MSSQL_SEASON_VIEWS": "Vision_FA26_Boxed_EUROPE"})
    assert digest.gather(db)["season"] == "FA26"


def test_newest_season_wins_when_two_views_are_configured():
    """Gecis doneminde eski gorunum listede kaliyor; ozet yeni sezonu yazmali."""
    config.apply_overrides({"MSSQL_SEASON_VIEWS":
                            "Vision_SP26_Boxed_EUROPE,Vision_FA26_Boxed_EUROPE"})
    assert config.season_label() == "FA26"


def test_next_years_season_beats_this_years():
    config.apply_overrides({"MSSQL_SEASON_VIEWS":
                            "Vision_FA26_Boxed_EUROPE,Vision_SP27_Boxed_EUROPE"})
    assert config.season_label() == "SP27"


def test_gather_honours_stale_days(db):
    """Ayni parca 90 gunluk esikte takilmis sayilmamali."""
    _parcel(db, status="in_transit")
    old = (datetime.now() - timedelta(days=30)).isoformat(timespec="seconds")
    db.conn.execute("UPDATE parcels SET handed_over_at = ?, last_event_at = ?, "
                    "created_at = ?", (old, old, old))
    db.conn.commit()
    assert len(digest.gather(db, stale_days=7)["stale"]) == 1
    assert digest.gather(db, stale_days=90)["stale"] == []


def test_gather_strips_trailing_slash_from_panel_url(db):
    assert digest.gather(db, panel_url="http://x:8765/")["panel_url"] == "http://x:8765"


# ------------------------------------------------------------------- render
def test_render_subject_is_english(db):
    """Ozet kurum disina gidiyor — panel Turkce olsa da mail Ingilizce (DOC_LANG)."""
    subject, _ = digest.render(digest.gather(db))
    assert "GLS" in subject and "daily summary" in subject
    assert "delivered" in subject and "problems" in subject


def test_subject_starts_with_system_name(db):
    """Gonderen adi Exchange'in elinde — maili kimin yazdigi KONUDAN anlasilmali."""
    subject, html = digest.render(digest.gather(db))
    assert subject.startswith(digest.SYSTEM_NAME)
    assert digest.SYSTEM_NAME in html


def test_mail_html_has_no_style_block(db):
    """Outlook <style> bloklarini yok sayar — her stil satir ici olmali."""
    _, html = digest.render(digest.gather(db))
    assert "<style" not in html
    assert "display:flex" not in html and "display:grid" not in html


def test_render_points_to_attachment(db):
    _parcel(db, status="exception")
    _, html = digest.render(digest.gather(db))
    assert "The attached Excel" in html


def test_render_states_season_totals(db):
    config.apply_overrides({"MSSQL_SEASON_VIEWS": "Vision_FA26_Boxed_EUROPE"})
    _parcel(db, tracking_no="38120177000001", status="delivered")
    _parcel(db, tracking_no="38120177000002", status="in_transit")
    subject, html = digest.render(digest.gather(db))
    assert "FA26" in subject and "FA26" in html
    assert "1 of 2 shipments" in html


def test_render_hides_panel_button_without_url(db):
    """Dugme metni i18n'den okunur: elle yazilan kopya panel adi degisince
    testi sessizce "hep gecer" hale getiriyordu."""
    _, html = digest.render(digest.gather(db))
    assert i18n.t("mail.open_panel") not in html


# -------------------------------------------------------------------- excel
def _workbook(db, **kw):
    import io

    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(excel.build(digest.gather(db, **kw))))


def test_excel_has_manual_list_headers(db):
    _parcel(db, status="exception")
    ws = _workbook(db).active
    header = [c.value for c in ws[1]]
    assert header[:3] == ["FLAG", "PARCEL NUMBER", "SHIPMENT AGENCY"]
    assert ws.cell(row=2, column=3).value == "Gls Netherlands"


def test_excel_lists_every_parcel_once(db):
    """Ek kumulatif: teslim edilmemis parcalar da listede olmali, tekrar olmadan."""
    _parcel(db, tracking_no="38120177000001", status="delivered")
    _parcel(db, tracking_no="38120177000002", status="created")
    ws = _workbook(db).active
    numbers = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert sorted(numbers) == ["38120177000001", "38120177000002"]


def test_excel_paints_delivered_rows_green(db):
    _parcel(db, tracking_no="38120177000001", status="delivered")
    _parcel(db, tracking_no="38120177000002", status="created")
    ws = _workbook(db).active
    filled = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=1).fill.patternType
              for r in range(2, ws.max_row + 1)}
    assert filled["38120177000001"] == "solid"
    assert filled["38120177000002"] is None


def test_excel_second_sheet_is_delivered_only(db):
    _parcel(db, tracking_no="38120177000001", status="delivered")
    _parcel(db, tracking_no="38120177000002", status="created")
    wb = _workbook(db)
    assert wb.sheetnames == ["All Parcels", "Delivered", "Problematic"]
    ws = wb["Delivered"]
    assert [ws.cell(row=r, column=2).value
            for r in range(2, ws.max_row + 1)] == ["38120177000001"]


def test_excel_third_sheet_collects_every_problem_kind(db):
    """Elle tutulan listedeki "Problematic Packages" sekmesinin karsiligi."""
    _parcel(db, tracking_no="38120177000001", status="exception")
    # "POD eksik" gun kapisi: teslim edileli >= stale_days GUN gecmis olmali
    # (bkz. problem_parcels, kullanici 2026-09-04).
    _parcel(db, tracking_no="38120177000002", status="delivered",
            delivered_date="2020-01-01T10:00:00Z")             # POD'u eksik
    _parcel(db, tracking_no="38120177000003", status="delivered",
            delivered_date="2020-01-01T10:00:00Z")
    db.set_pod_path("38120177000003", "/tmp/var.png", "png")
    ws = _workbook(db)["Problematic"]
    assert {ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)} == {
        "38120177000001", "38120177000002"}


def test_excel_problem_sheet_skips_parcels_gls_has_not_picked_up(db):
    """Kullanicinin karari: kamyonu bekleyen koli sorun degil, sayfayi sisiriyordu.

    Tam listede (ilk sayfa) FLAG sutunundan hala okunabiliyor.
    """
    _parcel(db, tracking_no="38120177000001", status="created",
            shipment_date="2020-01-01")
    wb = _workbook(db)
    assert wb["Problematic"].max_row == 1                    # yalnizca baslik satiri
    assert wb["All Parcels"].cell(row=2, column=1).value      # FLAG hala dolu


def test_excel_translates_the_delivery_note_we_wrote_ourselves(db):
    """GLS'in olay metinleri Ingilizce gelir; "teslim alan" satirini biz TR yaziyoruz.

    Dosya disariya gidiyor — EXPLAIN sutununda Turkce kalmamali.
    """
    _parcel(db, status="delivered")
    db.update_tracking_info("38120177000001",
                            last_event_text=delivered_note("Léonie"))
    ws = _workbook(db).active
    column = [c.value for c in ws[1]].index("EXPLAIN") + 1
    assert ws.cell(row=2, column=column).value == "Delivered, received by: Léonie"


def test_problem_listed_once_even_with_two_kinds():
    """Ayni parca hem hareketsiz hem POD'suz olabilir — tek satir yazilmali."""
    data = {"exceptions": [{"tracking_no": "X"}], "stale": [{"tracking_no": "X"}],
            "pod_missing": [{"tracking_no": "X"}], "partial": [{"tracking_no": "X"}],
            "damaged": [{"tracking_no": "X"}], "not_handed_over": [{"tracking_no": "X"}]}
    assert len(excel._problem_rows(data)) == 1


def test_excel_filename_carries_season(db):
    config.apply_overrides({"MSSQL_SEASON_VIEWS": "Vision_FA26_Boxed_EUROPE"})
    assert excel.filename(digest.gather(db)).startswith("FA26_Tracking_List_")


# --------------------------------------------------------------------- mail
def test_get_token_without_credentials_is_turkish(monkeypatch):
    monkeypatch.setattr(mail.requests, "post",
                        lambda *a, **kw: pytest.fail("aga cikilmamaliydi"))
    token, detail = mail.get_token()
    assert token is None
    assert "Client Secret" in detail


def test_send_refuses_empty_recipients(monkeypatch):
    monkeypatch.setattr(mail.requests, "post",
                        lambda *a, **kw: pytest.fail("aga cikilmamaliydi"))
    ok, detail = mail.send([], "konu", "<p>x</p>")
    assert ok is False and detail == "Alıcı listesi boş"


def test_send_is_inert_in_mock_mode(monkeypatch):
    monkeypatch.setattr(mail.requests, "post",
                        lambda *a, **kw: pytest.fail("mock modda aga cikilmamaliydi"))
    config.apply_overrides({"MAIL_SENDER": "a@b.c"})
    ok, detail = mail.send(["x@y.z"], "konu", "<p>x</p>")
    assert ok is True and "mock" in detail


def test_send_treats_202_as_success(monkeypatch):
    """Graph sendMail 202 doner, 200 DEGIL."""
    monkeypatch.setattr(config, "IS_MOCK", False)
    monkeypatch.setattr(mail, "get_token", lambda: ("tok", "ok"))
    config.apply_overrides({"MAIL_SENDER": "a@b.c"})
    monkeypatch.setattr(mail.requests, "post",
                        lambda *a, **kw: type("R", (), {"status_code": 202, "text": ""})())
    ok, _ = mail.send(["x@y.z"], "konu", "<p>x</p>")
    assert ok is True


def test_timeout_warns_that_mail_may_have_been_sent(monkeypatch):
    """Zaman asiminda sonuc BILINMEZ; "gonderilemedi" deyip tekrar denetmemeli."""
    monkeypatch.setattr(config, "IS_MOCK", False)
    monkeypatch.setattr(mail, "get_token", lambda: ("tok", "ok"))
    config.apply_overrides({"MAIL_SENDER": "a@b.c"})

    def _timeout(*a, **kw):
        raise mail.requests.Timeout("read timed out")

    monkeypatch.setattr(mail.requests, "post", _timeout)
    ok, detail = mail.send(["x@y.z"], "konu", "<p>x</p>")
    assert ok is False and "gitmiş olabilir" in detail


def test_from_field_is_not_sent(monkeypatch):
    """Exchange govdedeki gonderen adini yok sayar (canli denendi) — alan yollanmaz.

    Gelen kutusunda gorunen ad posta kutusunun M365'teki adidir; govdeye ad
    koymak 202 doner ama hicbir sey degistirmez. Yalan bir alan gonderilmesin.
    """
    monkeypatch.setattr(config, "IS_MOCK", False)
    monkeypatch.setattr(mail, "get_token", lambda: ("tok", "ok"))
    config.apply_overrides({"MAIL_SENDER": "a@b.c"})
    sent = {}

    def _post(url, json=None, **kw):
        sent.update(json)
        return type("R", (), {"status_code": 202, "text": ""})()

    monkeypatch.setattr(mail.requests, "post", _post)
    mail.send(["x@y.z"], "konu", "<p>x</p>")
    assert "from" not in sent["message"]


def test_send_rejects_oversized_attachments(monkeypatch):
    """Tavan tek eke degil, eklerin TOPLAMINA uygulanir."""
    monkeypatch.setattr(config, "IS_MOCK", False)
    monkeypatch.setattr(mail, "get_token", lambda: ("tok", "ok"))
    config.apply_overrides({"MAIL_SENDER": "a@b.c"})
    monkeypatch.setattr(mail.requests, "post",
                        lambda *a, **kw: pytest.fail("buyuk ek gonderilmemeliydi"))
    half = b"0" * (mail.MAX_ATTACHMENT_BYTES // 2 + 1)
    ok, detail = mail.send(["x@y.z"], "k", "<p>x</p>",
                           attachments=[("b.xlsx", half), ("p.zip", half)])
    assert ok is False and "çok büyük" in detail


def test_send_does_not_log_addresses(monkeypatch, caplog):
    config.apply_overrides({"MAIL_SENDER": "a@b.c"})
    with caplog.at_level("INFO"):
        mail.send(["gizli@alici.com"], "konu", "<p>x</p>")
    assert "gizli@alici.com" not in caplog.text


# ----------------------------------------------------------------------- run
def test_run_without_recipients_stays_offline(db, monkeypatch):
    monkeypatch.setattr(digest.mail, "send",
                        lambda *a, **kw: pytest.fail("gonderilmemeliydi"))
    ok, detail = digest.run(db, recipients=[])
    assert ok is False and detail == "Alıcı listesi boş"


def _spy_send(monkeypatch) -> dict:
    seen: dict = {}

    def fake_send(to, subject, html, attachments=None):
        seen.update(to=to, html=html, attachments=attachments or [])
        return True, "ok"

    monkeypatch.setattr(digest.mail, "send", fake_send)
    return seen


def test_run_attaches_excel_when_rows_exist(db, monkeypatch):
    _parcel(db, status="exception")
    seen = _spy_send(monkeypatch)
    digest.run(db, recipients=["a@b.c"])
    name, content = seen["attachments"][0]
    assert name.endswith(".xlsx") and content[:2] == b"PK"   # xlsx = zip


def test_run_sends_nothing_extra_when_db_empty(db, monkeypatch):
    seen = _spy_send(monkeypatch)
    digest.run(db, recipients=["a@b.c"])
    assert seen["attachments"] == []


# ------------------------------------------------------------------ POD zip
# Zip icindeki tam yol (bkz. web/routers/pod.py:pod_archive_path). Test parcasinin
# magaza kodu ve faturasi yok, sevkiyat sekli de bos — ad takip numarasina duser,
# klasorler de varsayilanlarina.
POD_IN_ZIP = "FA26_POD/Genel Sevkiyatlar/Bilinmiyor/38120177000001.pdf"


def _delivered_with_pod(db, path) -> None:
    # `pod_archive_path` sezon klasorunu `config.season_label()`den turetir
    # (bkz. web/routers/pod.py); baska testler zaten ayni sekilde ayarliyor.
    config.apply_overrides({"MSSQL_SEASON_VIEWS": "Vision_FA26_Boxed_EUROPE"})
    _parcel(db, status="delivered")
    db.set_pod_path("38120177000001", str(path), path.suffix.lstrip("."))
    _event(db, "38120177000001", "delivered", _now_iso(hours_ago=1))


def _png(path):
    from PIL import Image
    Image.new("RGB", (200, 100), "white").save(path, format="PNG")
    return path


def test_pod_zip_bundles_pdf_documents(db, tmp_path):
    """Ham goruntu ciplak bir fotograftir; eke panelin verdigi PDF belge girer."""
    _delivered_with_pod(db, _png(tmp_path / "pod.png"))

    data = digest.gather(db)
    name, content = digest.pod_zip(data)
    assert name.endswith(".zip") and data["pod_zip_count"] == 1
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        assert zf.namelist() == [POD_IN_ZIP]
        assert zf.read(POD_IN_ZIP).startswith(b"%PDF")


def test_ready_made_pdf_pod_is_not_rewrapped(db, tmp_path):
    """ShipIT/T&T zaten eksiksiz bir PDF verir — ikinci kez sarilmaz."""
    pod = tmp_path / "pod.pdf"
    pod.write_bytes(b"%PDF-1.4 hazir")
    _delivered_with_pod(db, pod)
    _, content = digest.pod_zip(digest.gather(db))
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        assert zf.read(POD_IN_ZIP) == b"%PDF-1.4 hazir"


def test_pod_zip_is_none_without_pods(db):
    _parcel(db, status="delivered")
    _event(db, "38120177000001", "delivered", _now_iso(hours_ago=1))
    assert digest.pod_zip(digest.gather(db)) is None


def test_pod_zip_survives_missing_file(db, tmp_path):
    """Arsivden elle silinmis bir POD ozetin tamamini dusurmemeli."""
    _delivered_with_pod(db, tmp_path / "yok.png")
    assert digest.pod_zip(digest.gather(db)) is None


def test_pod_zip_survives_broken_image(db, tmp_path):
    """Belge uretilemezse ham goruntuye DUSULMEZ; o parca atlanir."""
    pod = tmp_path / "bozuk.png"
    pod.write_bytes(b"\x89PNG bozuk")
    _delivered_with_pod(db, pod)
    assert digest.pod_zip(digest.gather(db)) is None


# ----------------------------------------------------------------- scheduler
def _in_loop(fn):
    """AsyncIOScheduler.start() CALISAN bir olay dongusu ister; testler senkron."""
    async def main():
        return fn()
    return asyncio.run(main())


def test_disabled_digest_adds_no_job(db):
    def body():
        d = DailyDigest(db=db, enabled=False)
        d.start()
        try:
            assert d.scheduler.get_job(JOB_ID) is None
        finally:
            d.shutdown()
    _in_loop(body)


def test_enabling_adds_job_and_hour_change_reschedules(db):
    def body():
        d = DailyDigest(db=db, enabled=False)
        d.start()
        try:
            d.reconfigure(enabled=True, hours=[6])
            assert d.scheduler.get_job(JOB_ID) is not None
            d.reconfigure(hours=[21])
            assert "21" in str(d.scheduler.get_job(JOB_ID).trigger)
            d.reconfigure(enabled=False)
            assert d.scheduler.get_job(JOB_ID) is None
        finally:
            d.shutdown()
    _in_loop(body)


def test_two_hours_become_one_cron_job(db):
    """Sabah + aksam TEK ise sigar: cron tetigi virgullu saat listesini kendi alir."""
    def body():
        d = DailyDigest(db=db, hours=[9, 18], enabled=True)
        d.start()
        try:
            assert "9,18" in str(d.scheduler.get_job(JOB_ID).trigger)
        finally:
            d.shutdown()
    _in_loop(body)


def test_hours_are_sorted_and_deduplicated(db):
    assert DailyDigest(db=db, hours=[18, 9, 9]).hours == [9, 18]


def test_a_broken_hour_list_never_silences_the_digest(db):
    """Ozetin hic gitmemesi fark edilmez; yanlis saatte gitmesi fark edilir."""
    assert DailyDigest(db=db, hours=[]).hours == [9]
    assert DailyDigest(db=db, hours=[25, 99]).hours == [9]


def test_hours_mean_turkish_time_not_the_container_clock(db):
    """Konteyner UTC calisiyor: saat dilimi verilmezse "sabah 9" 12:00'de giderdi."""
    def body():
        d = DailyDigest(db=db, hours=[9], enabled=True)
        d.start()
        try:
            trigger = d.scheduler.get_job(JOB_ID).trigger
            assert str(trigger.timezone) == "Europe/Istanbul"
        finally:
            d.shutdown()
    _in_loop(body)


def test_reconfigure_before_start_is_safe(db):
    """Zamanlayici baslamadan cagrilirsa patlamamali (CLI, testler)."""
    d = DailyDigest(db=db)
    d.reconfigure(enabled=True, hours=[9])
    assert d.hours == [9] and d.enabled is True


def test_send_now_records_last_run(db, monkeypatch):
    monkeypatch.setattr(digest, "run", lambda *a, **kw: (False, "Alıcı listesi boş"))
    d = DailyDigest(db=db)
    ok, detail = d.send_now()
    assert ok is False
    assert d.last_run_at and d.last_error == "Alıcı listesi boş"


# ---------------------------------------------------------------------------
# POD gelmedi, GLS'e sor
# ---------------------------------------------------------------------------

def _delivered_without_pod(db, tracking_no: str, days_ago: int) -> None:
    db.add_parcel(tracking_no=tracking_no, channel="NL", country="FR",
                  consignee_name="Emily Boutique", store_code="RZVFR2N",
                  invoice_number="5000044", status="delivered",
                  delivered_date=(datetime.now() - timedelta(days=days_ago))
                  .isoformat(timespec="seconds"))


def test_gls_is_not_asked_before_the_waiting_period(db):
    """POD teslimattan saatler sonra olusabiliyor; erken sorulursa GLS'e henuz
    var olmayan bir belge sorulmus olur."""
    _delivered_without_pod(db, "38120177000001", days_ago=2)
    assert db.pod_missing_since(wait_days=7) == []


def test_the_same_parcel_is_asked_only_once(db, monkeypatch):
    """Her sabah ayni numaralari yeniden yollamak GLS gozunde spam olur."""
    from notify import pod_chase

    sent = []
    monkeypatch.setattr(mail, "send",
                        lambda to, subject, html, **kw: (sent.append(to), (True, "ok"))[1])

    _delivered_without_pod(db, "38120177000002", days_ago=9)
    ok, _detail = pod_chase.run(db, wait_days=7, recipients=["gls@example.com"])
    assert ok and len(sent) == 1

    ok, detail = pod_chase.run(db, wait_days=7, recipients=["gls@example.com"])
    assert ok and len(sent) == 1 and "sorulacak parça yok" in detail


def test_a_failed_send_leaves_the_parcel_in_the_queue(db, monkeypatch):
    """Graph hatasi parcalari sessizce yutmamali — yarin yeniden denenmeli."""
    from notify import pod_chase

    monkeypatch.setattr(mail, "send", lambda *a, **kw: (False, "Graph 500"))
    _delivered_without_pod(db, "38120177000003", days_ago=9)
    ok, _ = pod_chase.run(db, wait_days=7, recipients=["gls@example.com"])
    assert ok is False
    assert len(db.pod_missing_since(wait_days=7)) == 1


def test_the_letter_repeats_what_the_user_wrote_by_hand(db):
    """Metin kullanicinin GLS'e elle yazdigi mektubun aynisi (2026-08-06)."""
    from notify import pod_chase

    _delivered_without_pod(db, "38120177005066", days_ago=9)
    subject, text, html_body = pod_chase.draft(db.pod_missing_since(wait_days=7))
    assert "Request for Official Proof of Delivery" in subject
    assert "Customer Number:" in text
    assert "Tracking No: 38120177005066 | Ref: RZVFR2N-5000044" in text
    assert "Recipient: Emily Boutique (FR)" in text
    # Mail HTML gider, panel duz metin gosterir — ikisi ayni satirlari tasir.
    assert "38120177005066" in html_body and "<li" in html_body


def test_the_inquiry_job_runs_even_when_the_digest_is_off(db):
    """Ayri istir: ozet kapaliyken de GLS'e sorulabilmeli."""
    from notify.scheduler import INQUIRY_JOB_ID

    def body():
        d = DailyDigest(db=db, enabled=False, inquiry_enabled=True)
        d.start()
        try:
            assert d.scheduler.get_job(INQUIRY_JOB_ID) is not None
            assert d.scheduler.get_job(JOB_ID) is None
            d.reconfigure(inquiry_enabled=False)
            assert d.scheduler.get_job(INQUIRY_JOB_ID) is None
        finally:
            d.shutdown()
    _in_loop(body)


# ------------------------------------------------ OpenWA oturum sagligi
# Canli olcum (2026-09-01, kullanici bildirdi "wp grubuna dun hicbir sey
# gelmedi"): oturum `qr_ready` durumundaydi — 29 Agustos'ta dusmus, UC GUN
# kimse fark etmemis, 31 Agustos'taki 23 teslimat bildirimi hic gitmemisti.
# OpenWA sunucusu AYAKTAYDI ve HTTP 200 donuyordu; sessizligi yalnizca oturum
# durumu ele veriyor.

def _wa_reset_cache():
    from notify import whatsapp as wa
    wa._status_cache.update(at=0.0, value=None)


def test_a_dropped_session_is_reported_as_not_ok(monkeypatch):
    from notify import whatsapp as wa
    from gls_api import config

    _wa_reset_cache()
    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    monkeypatch.setattr(config, "MODE", "prod")

    class R:
        status_code = 200
        def json(self): return {"status": "qr_ready"}
    monkeypatch.setattr(wa.requests, "get", lambda *a, **k: R())

    out = wa.session_status(force=True)
    assert out["ok"] is False
    assert out["status"] == "qr_ready"


def test_a_connected_session_is_ok(monkeypatch):
    from notify import whatsapp as wa
    from gls_api import config

    _wa_reset_cache()
    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    monkeypatch.setattr(config, "MODE", "prod")

    class R:
        status_code = 200
        def json(self): return {"status": "connected"}
    monkeypatch.setattr(wa.requests, "get", lambda *a, **k: R())

    assert wa.session_status(force=True)["ok"] is True


def test_an_unreachable_server_is_also_not_ok(monkeypatch):
    """Sunucuya ulasilamamak da "mesaj gitmiyor" demektir — operator gormeli."""
    from notify import whatsapp as wa
    from gls_api import config

    _wa_reset_cache()
    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    monkeypatch.setattr(config, "MODE", "prod")
    monkeypatch.setattr(wa.requests, "get",
                        lambda *a, **k: (_ for _ in ()).throw(OSError("baglanti yok")))

    out = wa.session_status(force=True)
    assert out["ok"] is False
    assert out["status"] == "unreachable"


def test_mock_mode_never_touches_the_network(monkeypatch):
    from notify import whatsapp as wa
    from gls_api import config

    _wa_reset_cache()
    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    monkeypatch.setattr(config, "MODE", "mock")
    monkeypatch.setattr(wa.requests, "get",
                        lambda *a, **k: pytest.fail("mock modda aga cikilmamaliydi"))

    assert wa.session_status(force=True)["ok"] is True


def test_the_status_is_cached_between_calls(monkeypatch):
    """Her sayfa acilisinda OpenWA'ya gidilmemeli."""
    from notify import whatsapp as wa
    from gls_api import config

    _wa_reset_cache()
    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    monkeypatch.setattr(config, "MODE", "prod")
    calls = []

    class R:
        status_code = 200
        def json(self): return {"status": "connected"}
    monkeypatch.setattr(wa.requests, "get", lambda *a, **k: calls.append(1) or R())

    wa.session_status(force=True)
    wa.session_status()
    wa.session_status()
    assert len(calls) == 1


def test_the_banner_appears_only_when_the_session_is_down(monkeypatch):
    from web import deps
    from notify import whatsapp as wa

    monkeypatch.setattr(wa, "session_status", lambda: {"ok": False, "status": "qr_ready", "detail": ""})
    assert deps.whatsapp_alarm()["status"] == "qr_ready"

    monkeypatch.setattr(wa, "session_status", lambda: {"ok": True, "status": "connected", "detail": ""})
    assert deps.whatsapp_alarm() is None


def test_a_broken_check_never_breaks_the_panel(monkeypatch):
    """Kontrolun kendisi patlarsa paneli dusurmemeli."""
    from web import deps
    from notify import whatsapp as wa

    monkeypatch.setattr(wa, "session_status",
                        lambda: (_ for _ in ()).throw(RuntimeError("bozuk")))
    assert deps.whatsapp_alarm() is None


# ------------------------------------------------- FedEx AYRI WhatsApp grubu
# Kullanici istegi (2026-09-01): FedEx kolileri ana GLS grubuna girmez ama
# KENDI grubuna bildirir; PDF olustugu anda EK ile, sonra hareket geldikce
# durum mesajiyla. Sozlesme canli yoklamayla cikarildi (alan adlari kucuk harf).

def test_the_document_payload_uses_the_lowercase_field_names(monkeypatch):
    from notify import whatsapp as wa
    from gls_api import config

    monkeypatch.setattr(config, "MODE", "prod")
    gonderilen = {}

    class R:
        status_code = 201
        text = '{"messageId":"x"}'
    monkeypatch.setattr(wa.requests, "post",
                        lambda url, **kw: gonderilen.update(url=url, **kw) or R())

    ok, _ = wa.send_document(b"%PDF-1.4", "876572877080.pdf", caption="merhaba",
                             chat_id="100000000000002@g.us")

    assert ok
    body = gonderilen["json"]
    assert set(body) == {"chatId", "base64", "mimetype", "filename", "caption"}
    assert body["mimetype"] == "application/pdf"
    assert body["chatId"] == "100000000000002@g.us"
    assert "send-document" in gonderilen["url"]


def test_the_created_message_carries_the_number_and_the_tracking_link(monkeypatch):
    from notify import whatsapp as wa

    yakalanan = {}
    monkeypatch.setattr(wa, "send_document",
                        lambda pdf, fn, caption="", chat_id="", **kw:
                            yakalanan.update(caption=caption, fn=fn, chat=chat_id) or (True, ""))

    wa.notify_fedex_created({"tracking_no": "876572877080", "store_code": "CA1"}, b"%PDF")

    assert "876572877080" in yakalanan["caption"]
    # Metin INGILIZCE (kullanici karari 2026-09-01).
    assert "has been created" in yakalanan["caption"]
    assert "fedex.com/fedextrack/?trknbr=876572877080" in yakalanan["caption"]
    # Ek adi panelden indirilenle AYNI (kullanici karari 2026-09-02).
    assert yakalanan["fn"] == "Shipment Details 876572877080.pdf"


def test_the_created_message_uses_the_filename_it_is_given(monkeypatch):
    """Ad cagirandan gelir (`tracking.box_image_filename`) — WhatsApp'a dusen
    dosya ile panelden inen dosya ayni adi tasisin."""
    from notify import whatsapp as wa

    yakalanan = {}
    monkeypatch.setattr(wa, "send_document",
                        lambda pdf, fn, caption="", chat_id="", **kw:
                            yakalanan.update(fn=fn) or (True, ""))

    wa.notify_fedex_created({"tracking_no": "876572877080"}, b"%PDF",
                            filename="Shipment Details 876572877080.pdf")
    assert yakalanan["fn"] == "Shipment Details 876572877080.pdf"


def test_the_created_message_falls_back_to_text_without_a_pdf(monkeypatch):
    """Belge uretilemediyse bildirim tumden kaybolmasin — metin olarak gider."""
    from notify import whatsapp as wa

    yakalanan = {}
    monkeypatch.setattr(wa, "send_text",
                        lambda text, chat_id="", **kw:
                            yakalanan.update(text=text) or (True, ""))
    monkeypatch.setattr(wa, "send_document",
                        lambda *a, **kw: pytest.fail("PDF yokken belge gonderilmemeli"))

    ok, _ = wa.notify_fedex_created({"tracking_no": "876572877080"})
    assert ok
    assert "876572877080" in yakalanan["text"]


def test_the_status_message_carries_the_tracking_link(monkeypatch):
    from notify import whatsapp as wa

    yakalanan = {}
    monkeypatch.setattr(wa, "send_text",
                        lambda text, chat_id="", **kw:
                            yakalanan.update(text=text, chat=chat_id) or (True, ""))

    wa.notify_fedex_status({"tracking_no": "876572877080", "status": "in_transit",
                            "last_event_text": "Picked up"})

    assert "876572877080" in yakalanan["text"]
    assert "fedex.com/fedextrack" in yakalanan["text"]
    assert "Picked up" in yakalanan["text"]


def test_fedex_events_never_reach_the_main_gls_group(monkeypatch, tmp_path):
    """Kullanicinin 2026-08-31 karari korunur: ana gruba SIZMAZ."""
    from notify import whatsapp as wa
    from gls_api import config
    from tracking.db import ShipmentsDB

    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    db = ShipmentsDB(path=tmp_path / "wa.db")
    db.add_parcel("876572877080", "FEDEX", consignee_name="c", country="CA",
                  status="in_transit")

    monkeypatch.setattr(wa, "notify_delivery",
                        lambda *a, **k: pytest.fail("ana gruba gitmemeliydi"))
    monkeypatch.setattr(wa, "notify_problem",
                        lambda *a, **k: pytest.fail("ana gruba gitmemeliydi"))
    cagrilar = []
    monkeypatch.setattr(wa, "notify_fedex_status",
                        lambda p, **k: cagrilar.append(p["tracking_no"]) or (True, ""))

    wa.handle_tracker_event(db, {"tracking_no": "876572877080",
                                 "channel": "FEDEX", "new_status": "in_transit"})

    assert cagrilar == ["876572877080"]


def test_created_status_is_not_announced_twice(monkeypatch, tmp_path):
    """"Olusturuldu" mesaji PDF ile EK olarak gidiyor; tarayici tekrar duyurmasin."""
    from notify import whatsapp as wa
    from gls_api import config
    from tracking.db import ShipmentsDB

    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    db = ShipmentsDB(path=tmp_path / "wa2.db")
    db.add_parcel("876572877080", "FEDEX", consignee_name="c", country="CA",
                  status="created")
    monkeypatch.setattr(wa, "notify_fedex_status",
                        lambda *a, **k: pytest.fail("created icin mesaj gitmemeliydi"))

    wa.handle_tracker_event(db, {"tracking_no": "876572877080",
                                 "channel": "FEDEX", "new_status": "created"})


def test_nothing_is_sent_when_the_fedex_group_is_not_configured(monkeypatch, tmp_path):
    from notify import whatsapp as wa
    from gls_api import config
    from tracking.db import ShipmentsDB

    monkeypatch.setattr(config, "OPENWA_ENABLED", "1")
    monkeypatch.setattr(config, "OPENWA_FEDEX_CHAT_ID", "")
    db = ShipmentsDB(path=tmp_path / "wa3.db")
    db.add_parcel("876572877080", "FEDEX", consignee_name="c", country="CA",
                  status="delivered")
    monkeypatch.setattr(wa, "notify_fedex_status",
                        lambda *a, **k: pytest.fail("grup yokken gitmemeliydi"))

    wa.handle_tracker_event(db, {"tracking_no": "876572877080",
                                 "channel": "FEDEX", "new_status": "delivered"})


def test_fedex_messages_are_always_english(monkeypatch):
    """Kullanici karari (2026-09-01): FedEx bildirimlerinin TAMAMI Ingilizce.

    `status_label(code)` argumansiz cagrilirsa ISTEGIN dilini alir; arka plan
    tarayicisinda bu Turkce'ye duser ve gruba "Yolda"/"Teslim Edildi" giderdi.
    """
    from notify import whatsapp as wa

    tut = {}
    monkeypatch.setattr(wa, "send_text",
                        lambda t, chat_id="", **k: tut.update(t=t) or (True, ""))

    for status, beklenen in (("in_transit", "In Transit"),
                             ("out_for_delivery", "Out for Delivery"),
                             ("delivered", "Delivered"),
                             ("returned", "Returned")):
        wa.notify_fedex_status({"tracking_no": "876572877080", "status": status})
        assert f"Status: *{beklenen}*" in tut["t"], status
        for turkce in ("Yolda", "Teslim", "Dağıtımda", "İade", "Durum:"):
            assert turkce not in tut["t"]


def test_the_created_message_leaves_the_box_contents_out(monkeypatch):
    """Kullanici karari (2026-09-02): "bunu mesaja ekleme".

    Icerik listesi gercek veride coguna "— 1" yazan 15+ satir uretiyordu;
    ayrinti zaten ekteki PDF'te.
    """
    from notify import whatsapp as wa

    tut = {}
    monkeypatch.setattr(wa, "send_document",
                        lambda pdf, fn, caption="", chat_id="", **k:
                            tut.update(c=caption) or (True, ""))

    wa.notify_fedex_created({"tracking_no": "T1", "store_code": "CA1"}, b"%PDF")

    assert "Box contents" not in tut["c"]
    # Alici adi da yok — bos alanda "for *Müşteri*" gibi Turkce bir doldurma
    # metni cikiyordu (kullanici karari 2026-09-02: "oluşturuldu demen yeterli").
    assert "for" not in tut["c"]
    assert "Müşteri" not in tut["c"]
    assert "CA1" not in tut["c"]
    assert "•" not in tut["c"]
    assert tut["c"].splitlines()[-1] == "fedex.com/fedextrack/?trknbr=T1"


def test_the_tracking_link_is_short_enough_for_one_line():
    """Kullanici istegi (2026-09-02): "bu kadar kocaman olmasin baglanti linki".

    Uzun `https://www.fedex.com/...` biçimi WhatsApp balonunda ÜÇ SATIRA
    sariyordu. Ciplak alan adi da tiklanabilir kalir.
    """
    from notify.whatsapp import fedex_tracking_url

    url = fedex_tracking_url("876183992148")
    assert url == "fedex.com/fedextrack/?trknbr=876183992148"
    assert len(url) <= 45
    assert not url.startswith("http")


def test_the_status_message_asks_whatsapp_not_to_draw_a_preview_card(monkeypatch):
    """DEV onizleme karti (kapak gorseli + baslik) mesaji ezip geciyordu."""
    from notify import whatsapp as wa

    tut = {}
    monkeypatch.setattr(wa, "send_text",
                        lambda t, chat_id="", **k: tut.update(k) or (True, ""))
    wa.notify_fedex_status({"tracking_no": "T1", "status": "in_transit"})
    assert tut["link_preview"] is False


def test_send_text_forwards_the_preview_flag_to_the_gateway(monkeypatch):
    """Alan adi CANLI dogrulandi: `linkPreview` 201, `linkpreview`/`preview` 400."""
    from notify import whatsapp as wa

    gonderilen = {}

    class _Yanit:
        status_code, text = 201, "{}"

    monkeypatch.setattr(wa.config, "MODE", "live", raising=False)
    monkeypatch.setattr(wa.requests, "post",
                        lambda url, json=None, headers=None, timeout=None:
                            gonderilen.update(json) or _Yanit())

    wa.send_text("x", chat_id="c@g.us", link_preview=False)
    assert gonderilen["linkPreview"] is False

    wa.send_text("x", chat_id="c@g.us")
    assert gonderilen["linkPreview"] is True


def test_the_daily_summary_lists_shipments_from_the_last_seven_days(monkeypatch):
    """Kullanici istegi (2026-09-02): "aksam 6 da ... son 7 gonderi statuleri
    ve linklerinin oldugu TEK bir mesaj ... son 7 günü baz alacağı için orada
    last 7 days yazabilir" — sabit sayida kayit DEGIL, TARIH penceresi."""
    from datetime import datetime
    from notify import whatsapp as wa

    monkeypatch.setattr(wa.timez, "now", lambda: datetime(2026, 9, 2, 18, 0))

    class SahteDB:
        def list_parcels(self, channel=None, limit=500):
            assert channel == "FEDEX"
            return [
                {"tracking_no": "IN", "status": "in_transit",
                 "shipment_date": "2026-08-27"},           # tam sinirda: 6 gun once
                {"tracking_no": "OUT", "status": "delivered",
                 "shipment_date": "2026-08-20"},           # 13 gun once: PENCERE DISI
            ]

    tut = {}
    monkeypatch.setattr(wa, "send_text",
                        lambda t, chat_id="", **k: tut.update(t=t, k=k) or (True, ""))

    ok, _ = wa.notify_fedex_daily_summary(SahteDB())

    assert ok is True
    assert "*IN*" in tut["t"]
    assert "*OUT*" not in tut["t"]
    assert "fedex.com/fedextrack" in tut["t"]
    # Baslikta artik degisken bir sayi yok — o gunku koli sayisina gore
    # degisiyordu (kullanici karari: "yazmasına gerek yok").
    assert "Daily Summary (Last 7 Days)" in tut["t"]
    assert tut["k"]["link_preview"] is False


def test_the_daily_summary_says_nothing_when_there_is_no_fedex_parcel(monkeypatch):
    from notify import whatsapp as wa

    class BosDB:
        def list_parcels(self, channel=None, limit=500):
            return []

    tut = {"called": False}
    monkeypatch.setattr(wa, "send_text", lambda *a, **k: tut.update(called=True))

    ok, detail = wa.notify_fedex_daily_summary(BosDB())
    assert ok is True
    assert tut["called"] is False


def test_the_created_message_carries_the_shipment_date(monkeypatch):
    """Kullanici istegi (2026-09-02): "gönderi tarihini de ekleyelim", sonra
    "olmasına gerek yok [ayrı 'Created at:' satırı] ... has been created at
    24.08.2026. olması daha makul" — TEK cumleye gomulu, ayri satir DEGIL."""
    from notify import whatsapp as wa

    tut = {}
    monkeypatch.setattr(wa, "send_document",
                        lambda pdf, fn, caption="", chat_id="", **k:
                            tut.update(c=caption) or (True, ""))
    wa.notify_fedex_created({"tracking_no": "T1", "shipment_date": "2026-08-30"}, b"%PDF")
    assert "has been created at 30.08.2026." in tut["c"]
    assert "Created at:" not in tut["c"]


def test_the_status_message_carries_the_shipment_date(monkeypatch):
    from notify import whatsapp as wa

    tut = {}
    monkeypatch.setattr(wa, "send_text",
                        lambda t, chat_id="", **k: tut.update(t=t) or (True, ""))
    wa.notify_fedex_status({"tracking_no": "T1", "status": "delivered",
                            "shipment_date": "2026-08-30"})
    assert "Created at: 30.08.2026" in tut["t"]


def test_problem_and_returned_messages_carry_the_shipment_date(monkeypatch):
    from notify import whatsapp as wa

    tut = {}
    monkeypatch.setattr(wa, "send_text",
                        lambda t, chat_id="", **k: tut.update(t=t) or (True, ""))

    wa.notify_problem({"tracking_no": "T1", "shipment_date": "2026-08-30"})
    assert "30.08.2026" in tut["t"]

    wa.notify_returned({"tracking_no": "T1", "shipment_date": "2026-08-30"})
    assert "30.08.2026" in tut["t"]


def test_fedex_daily_summary_job_needs_the_group_and_whatsapp_enabled(db, monkeypatch):
    """Kullanici istegi (2026-09-02): "aksam 6 da ... son 7 gonderi" ozeti.

    FedEx grubu (`OPENWA_FEDEX_CHAT_ID`) tanimli degilse ya da WhatsApp
    kapaliysa is HIC eklenmez — `partial_deliveries_wa` ile ayni kural.
    """
    from gls_api import config as gls_config
    from notify.scheduler import FEDEX_SUMMARY_JOB_ID, DailyDigest

    def body():
        monkeypatch.setattr(gls_config, "OPENWA_ENABLED", "0", raising=False)
        monkeypatch.setattr(gls_config, "OPENWA_FEDEX_CHAT_ID", "", raising=False)
        d = DailyDigest(db=db)
        d.start()
        try:
            assert d.scheduler.get_job(FEDEX_SUMMARY_JOB_ID) is None

            monkeypatch.setattr(gls_config, "OPENWA_ENABLED", "1", raising=False)
            d.reconfigure()
            assert d.scheduler.get_job(FEDEX_SUMMARY_JOB_ID) is None  # grup yok

            monkeypatch.setattr(gls_config, "OPENWA_FEDEX_CHAT_ID", "120@g.us", raising=False)
            d.reconfigure()
            job = d.scheduler.get_job(FEDEX_SUMMARY_JOB_ID)
            assert job is not None
            assert "18" in str(job.trigger)
            assert str(job.trigger.timezone) == "Europe/Istanbul"
        finally:
            d.shutdown()
    _in_loop(body)


def test_send_fedex_summary_now_calls_the_whatsapp_helper(db, monkeypatch):
    from notify import whatsapp as wa
    from notify.scheduler import DailyDigest

    tut = {}
    monkeypatch.setattr(wa, "notify_fedex_daily_summary",
                        lambda db, **k: tut.update(called=True) or (True, "ok"))
    d = DailyDigest(db=db)
    ok, detail = d.send_fedex_summary_now()
    assert ok is True and tut["called"] is True
