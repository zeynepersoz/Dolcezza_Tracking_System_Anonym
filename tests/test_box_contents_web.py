# -*- coding: utf-8 -*-
"""`/tracking/contents/{no}` — koli icerigi ucu.

Bu ucun tek katı kurali var: HTMX bir 4xx/5xx yanitta hedefin govdesini
DEGISTIRMEZ. Hata durumunda kullanici acilan bolumde "Yükleniyor…" yazisiyla
kalir ve neyin bozuk oldugunu asla ogrenemez. Bu yuzden ERP kapali, kayit yok
ve baglanti coktu senaryolarinin UCU DE 200 doner; ne oldugunu govde soyler.
"""
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import i18n
from erp import contents as erp_contents
from erp.client import ERPError
from gls_api import config
from tracking.db import ShipmentsDB
from web import deps
from web.routers import tracking as tracking_router


@pytest.fixture
def db(tmp_path, monkeypatch):
    d = ShipmentsDB(path=tmp_path / "shipments.db")
    d.add_parcel(tracking_no="38120177016222", channel="NL", status="delivered",
                 consignee_name="AHA Austria", country="AT", season="FA26",
                 store_code="NUNAT1N")
    d.add_parcel(tracking_no="38120177016223", channel="NL", status="in_transit",
                 consignee_name="AHA Austria", country="AT", season="FA26",
                 store_code="NUNAT1N")
    monkeypatch.setattr(deps, "_shipments_db", d)
    return d


@pytest.fixture
def client(db):
    app = FastAPI()
    app.include_router(tracking_router.router)
    return TestClient(app)


@pytest.fixture
def erp_on(monkeypatch):
    """MSSQL yapilandirilmis gibi davran — testler ag baglantisi kurmaz."""
    monkeypatch.setattr(config, "mssql_configured", lambda: True)
    monkeypatch.setattr(erp_contents, "fetch_box_code", lambda *a, **k: "8682500093454")


def test_a_missing_parcel_is_a_404(client):
    """Panelde olmayan bir numaraya icerik sorulmaz — bu bir hata, durum degil."""
    assert client.get("/tracking/contents/38120177099999").status_code == 404


def test_the_matrix_is_drawn_when_the_erp_answers(client, erp_on, monkeypatch):
    monkeypatch.setattr(erp_contents, "fetch_contents", lambda *a, **k: [
        {"style": "76604", "color": "A/S", "size": "M", "qty": 5, "barcode": "111"},
    ])

    r = client.get("/tracking/contents/38120177016222")

    assert r.status_code == 200
    assert "76604" in r.text and "A/S" in r.text
    assert "8682500093454" in r.text  # koli barkodu basligi


def test_an_unconfigured_erp_still_answers_200(client, monkeypatch):
    monkeypatch.setattr(config, "mssql_configured", lambda: False)

    r = client.get("/tracking/contents/38120177016222")

    assert r.status_code == 200
    assert i18n.t("tracking.detail.contents_off", "tr") in r.text


def test_a_box_with_no_record_says_so_instead_of_showing_nothing(client, erp_on,
                                                                 monkeypatch):
    """FA26'nin 52 kolisi yalnizca `Erp_Box`'ta; gorunumlerde icerigi yok."""
    monkeypatch.setattr(erp_contents, "fetch_contents", lambda *a, **k: [])

    r = client.get("/tracking/contents/38120177016222")

    assert r.status_code == 200
    assert i18n.t("tracking.detail.contents_empty", "tr") in r.text


def test_the_box_can_be_taken_away_as_a_spreadsheet(client, erp_on, monkeypatch):
    """Depo bu tabloyu Excel'de suzuyor — model ve renk AYRI kolonlar."""
    monkeypatch.setattr(erp_contents, "fetch_contents", lambda *a, **k: [
        {"style": "76604", "color": "A/S", "size": "M", "qty": 5, "barcode": "111"},
    ])

    r = client.get("/tracking/contents/38120177016222/export.xlsx")

    assert r.status_code == 200
    assert "38120177016222.xlsx" in r.headers["content-disposition"]
    assert r.content[:2] == b"PK"  # xlsx bir zip'tir


def test_the_box_can_be_taken_away_as_a_pdf(client, erp_on, monkeypatch):
    monkeypatch.setattr(erp_contents, "fetch_contents", lambda *a, **k: [
        {"style": "76604", "color": "A/S", "size": "M", "qty": 5, "barcode": "111"},
    ])

    r = client.get("/tracking/contents/38120177016222/export.pdf")

    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"


def test_an_empty_box_produces_no_file_at_all(client, erp_on, monkeypatch):
    """Bos bir Excel "koli bos" der; oysa sorun ERP tarafinda olabilir."""
    monkeypatch.setattr(erp_contents, "fetch_contents", lambda *a, **k: [])

    assert client.get(
        "/tracking/contents/38120177016222/export.xlsx").status_code == 404


def test_a_broken_connection_answers_200_with_a_short_message(client, erp_on,
                                                              monkeypatch):
    """Ayrinti log'a gider: MSSQL hatalari sunucu ve kullanici adi icerebiliyor."""
    def boom(*a, **k):
        raise ERPError("192.0.2.30:1434 -> login failed for user 'sa'")

    monkeypatch.setattr(erp_contents, "fetch_contents", boom)

    r = client.get("/tracking/contents/38120177016222")

    assert r.status_code == 200
    assert i18n.t("tracking.detail.contents_error", "tr") in r.text
    assert "192.0.2.30" not in r.text


def test_only_one_number_stands_next_to_the_piece_count(client, erp_on, monkeypatch):
    """Ciplak "28" ile ciplak "34" ayni satirin iki ucundaydi ve esit sanildi.

    Barkod sayisi alt satirdan KALDIRILDI: iki rakam yan yana durdukca
    karsilastiriliyor. Kalan tek sayi model adedi, sagdaki adet birimli.
    """
    monkeypatch.setattr(erp_contents, "fetch_contents", lambda *a, **k: [
        {"style": "76604", "color": "A/S", "size": "M", "qty": 5, "barcode": "111"},
        {"style": "76604", "color": "A/S", "size": "L", "qty": 3, "barcode": "222"},
    ])

    text = client.get("/tracking/contents/38120177016222").text

    assert i18n.t("tracking.detail.contents_summary", "tr", models=1) in text
    assert i18n.t("export.pieces", "tr").lower() in text.lower()


def test_the_single_box_file_ends_with_a_total_line(client, erp_on, monkeypatch):
    """Kullanici "exelde de sikinti cikariyor" dedi: dosya kendi kendini
    dogrulamiyordu, toplam satiri yoktu."""
    monkeypatch.setattr(erp_contents, "fetch_contents", lambda *a, **k: [
        {"style": "76604", "color": "A/S", "size": "M", "qty": 5, "barcode": "111"},
        {"style": "76604", "color": "A/S", "size": "L", "qty": 3, "barcode": "222"},
    ])

    sheet = load_workbook(io.BytesIO(client.get(
        "/tracking/contents/38120177016222/export.xlsx").content)).active
    last = [c.value for c in sheet[sheet.max_row]]

    assert last[0] == i18n.t("export.total_row", "tr")
    assert last[-1] == 8


# ------------------------------------------------- toplu (musteri bazli) Excel

BULK = "/tracking/contents/export.xlsx"


@pytest.fixture
def erp_bulk(monkeypatch):
    """Yalnizca BIR kolinin icerigi cikar; digeri "kayit yok" olmali."""
    monkeypatch.setattr(config, "mssql_configured", lambda: True)
    monkeypatch.setattr(erp_contents, "fetch_many", lambda *a, **k: {
        "38120177016222": [
            {"style": "76604", "color": "A/S", "size": "M", "qty": 5,
             "barcode": "111", "box_no": "30894"},
        ],
    })
    monkeypatch.setattr(erp_contents, "fetch_box_codes", lambda *a, **k: {
        "38120177016222": "8682500093454"})


def _book(response):
    return load_workbook(io.BytesIO(response.content))


def test_the_selection_comes_down_as_one_file_with_three_sheets(client, erp_bulk):
    """Magazanin 40 kolisini isteyen kisi 40 dosya indirip elle birlestiriyordu."""
    r = client.post(BULK, data={"tracking_no": ["38120177016222",
                                                "38120177016223"]})

    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    assert len(_book(r).worksheets) == 3


def test_the_file_is_named_after_the_store_when_there_is_only_one(client, erp_bulk):
    """Kullanicinin derdi zaten "NUNAT1N'in kolileri"."""
    r = client.post(BULK, data={"tracking_no": ["38120177016222"]})

    assert "NUNAT1N" in r.headers["content-disposition"]


def test_a_box_without_contents_is_still_listed_with_a_note(client, erp_bulk):
    """Sessizce atlanirsa kullanici o koliyi "bos" sanir — canlida 71 koli boyle."""
    summary = _book(client.post(BULK, data={
        "tracking_no": ["38120177016222", "38120177016223"]})).worksheets[0]
    rows = [[c.value for c in row] for row in summary.iter_rows(min_row=2)]

    absent = next(r for r in rows if r[0] == "38120177016223")
    assert absent[-1] == i18n.t("export.no_records", "tr")


def test_the_summary_ends_with_a_total_line(client, erp_bulk):
    summary = _book(client.post(BULK, data={
        "tracking_no": ["38120177016222", "38120177016223"]})).worksheets[0]

    assert summary[summary.max_row][0].value == i18n.t("export.total_row", "tr")


def test_a_made_up_tracking_number_never_reaches_the_erp(client, erp_bulk):
    """Numaralar formdan geliyor: panelde OLMAYAN bir numara ERP'ye sorulmaz."""
    summary = _book(client.post(BULK, data={
        "tracking_no": ["38120177016222", "38120177099999"]})).worksheets[0]
    numbers = {row[0].value for row in summary.iter_rows(min_row=2)}

    assert "38120177099999" not in numbers


def test_an_empty_selection_is_a_400(client, erp_bulk):
    """Bu bir DOSYA INDIRME, HTMX degil — gercek hata kodu kullanilabilir."""
    assert client.post(BULK, data={"tracking_no": []}).status_code == 400


def test_too_many_boxes_are_refused_instead_of_hanging_the_erp(client, erp_bulk,
                                                               monkeypatch):
    monkeypatch.setattr(tracking_router, "CONTENTS_BULK_LIMIT", 1)

    r = client.post(BULK, data={"tracking_no": ["38120177016222",
                                                "38120177016223"]})

    assert r.status_code == 400


def test_an_unconfigured_erp_says_so_instead_of_sending_an_empty_book(client,
                                                                      monkeypatch):
    monkeypatch.setattr(config, "mssql_configured", lambda: False)

    r = client.post(BULK, data={"tracking_no": ["38120177016222"]})

    assert r.status_code == 503


def test_the_bulk_excel_has_autofilter_freeze_and_pastel_colors(client, erp_bulk):
    """Kullanici talebi: ust satirda filtre + dondurme, 2. sayfada ayni koli pastel renk."""
    r = client.post(BULK, data={"tracking_no": ["38120177016222"]})
    book = _book(r)

    for sheet in book.worksheets:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None

    # 2. sayfa (Icerik)
    lines_sheet = book.worksheets[1]
    assert lines_sheet.max_row >= 2
    data_cell = lines_sheet.cell(row=2, column=1)
    assert data_cell.fill is not None
    assert data_cell.fill.fgColor.rgb is not None

