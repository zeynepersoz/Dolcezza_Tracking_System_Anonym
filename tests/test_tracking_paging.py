# -*- coding: utf-8 -*-
"""Takip listesinin sayfalanmasi.

Once tum liste (1000 satirlik tavan) tek seferde ciziliyordu ve tablo kendini
30 saniyede bir yeniliyordu — her tur 1000 satirlik HTML uretip DOM'a yikmak.
Kullanici bunu "memoryi sisirir ve listeler cok uzun oldugu icin yuklenmesi cok
uzun surer" diye bildirdi. Artik ekrana bir SAYFA iner (`PAGE_SIZE`), oklarla
gezilir; dosyaya ise suzgece uyan HER satir gider.
"""
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from tracking.db import ShipmentsDB
from web import deps
from web.routers import tracking as tracking_router

SIZE = tracking_router.PAGE_SIZE


@pytest.fixture
def db(tmp_path, monkeypatch):
    d = ShipmentsDB(path=tmp_path / "shipments.db")
    # Iki tam sayfa + 5 satir: son sayfanin yarim oldugu durum da denensin.
    for i in range(2 * SIZE + 5):
        d.add_parcel(tracking_no=f"381201770{i:05d}", channel="NL",
                     status="delivered" if i % 2 else "in_transit",
                     consignee_name=f"Alici {i}", country="AT")
    monkeypatch.setattr(deps, "_shipments_db", d)
    return d


@pytest.fixture
def client(db):
    app = FastAPI()
    app.include_router(tracking_router.router)
    return TestClient(app)


# ---------------------------------------------------------------- veri katmani
def test_a_page_is_a_window_not_the_whole_list(db):
    assert len(db.list_parcels(limit=SIZE)) == SIZE


def test_the_next_page_does_not_repeat_the_first(db):
    """`updated_at` toplu senkronda YUZLERCE satirda ayni olur.

    Esitlerin sirasi SQLite'in keyfine kalsaydi 2. sayfa 1. sayfadaki satirlari
    tekrar getirir, baskalari hic gorunmezdi — siralamanin ikinci anahtari var.
    """
    first = {r["tracking_no"] for r in db.list_parcels(limit=SIZE)}
    second = {r["tracking_no"] for r in db.list_parcels(limit=SIZE, offset=SIZE)}

    assert len(first) == len(second) == SIZE
    assert not (first & second)


def test_the_last_page_is_short_not_padded(db):
    assert len(db.list_parcels(limit=SIZE, offset=2 * SIZE)) == 5


def test_the_counter_and_the_page_share_one_filter(db):
    """Ikisi ayri kurulsaydi "986 kayit" derken sayfalarda 900 satir cikardi."""
    total = db.count_parcels(status="delivered")
    rows = db.list_parcels(status="delivered", limit=10_000)

    assert total == len(rows) == SIZE + 2


def test_an_unknown_filter_field_is_refused_by_the_counter(db):
    """Sayac da liste ile ayni beyaz listeden gecmeli."""
    with pytest.raises(TypeError):
        db.count_parcels(uydurma_alan="x")


# ------------------------------------------------------------------------- uc
def test_the_screen_gets_one_page_not_the_whole_list(client):
    rows = client.get("/tracking/rows").text

    assert rows.count('class="row-pick"') == SIZE


def test_the_arrow_walks_to_the_next_page(client, db):
    first = client.get("/tracking/rows").text
    second = client.get("/tracking/rows?page=2").text

    newest = db.list_parcels(limit=1)[0]["tracking_no"]
    assert newest in first and newest not in second


def test_the_foot_says_where_we_are(client):
    rows = client.get("/tracking/rows?page=2").text

    assert f"{SIZE + 1}-{2 * SIZE} / {2 * SIZE + 5}" in rows


def test_a_made_up_page_number_does_not_crash_the_list(client):
    """Sayi istekten geliyor; ciplak `int()` `?page=abc`de 500 dondururdu."""
    assert client.get("/tracking/rows?page=abc").status_code == 200
    assert client.get("/tracking/rows?page=-4").status_code == 200


def test_a_page_past_the_end_falls_back_to_the_last_full_one(client):
    """Suzgec daralinca kullanici 7. sayfada kalabilir; bos tablo yerine
    son dolu sayfa gelir."""
    rows = client.get("/tracking/rows?page=99").text

    assert rows.count('class="row-pick"') == 5
    assert f"{2 * SIZE + 1}-{2 * SIZE + 5} / {2 * SIZE + 5}" in rows


def test_the_arrows_stay_home_when_everything_fits(client):
    """Tek sayfalik listede ok cizilmez — tiklanacak yeri olmayan iki dugme."""
    rows = client.get("/tracking/rows?f_consignee_name=Alici 7").text

    assert "goToPage" not in rows


def test_the_file_still_carries_every_matching_row_not_one_page(client):
    """Dosya EKSIK OLMAMALI: kullanici onu Excel'de suzecek ve orada eksik
    satir sessizce yanlis toplam demek."""
    from openpyxl import load_workbook
    import io

    sheet = load_workbook(io.BytesIO(
        client.get("/tracking/export.xlsx").content)).active

    assert sheet.max_row == 2 * SIZE + 5 + 1  # + baslik satiri


def test_clearing_the_filters_leaves_the_address_behind_too():
    """Panelden "Tumunu gor" ile gelindiginde adres `?delivered_day=2026-08-19`.

    Alanlari JS ile bosaltmak tabloyu tazeliyordu ama adres, yesil "Bugun teslim
    edilen" etiketi ve "3 filtre" rozeti SUNUCUDAN gelmisti — kullanici
    temizledigini sanip yenileyince filtre geri geliyordu.
    """
    from pathlib import Path

    body = Path("web/templates/tracking.html").read_text(encoding="utf-8")
    reset = body.split("function resetFilters()")[1].split("}")[0]

    assert "/tracking'" in reset


def test_the_page_number_rides_along_with_the_auto_refresh():
    """Tablo 30 saniyede bir kendini `hx-include` ile yeniliyor.

    `page` disarida kalsaydi 3. sayfaya bakan kullanici yarim dakika sonra
    sessizce 1. sayfaya donerdi.
    """
    from pathlib import Path

    page = Path("web/templates/tracking.html").read_text(encoding="utf-8")
    include = page.split('hx-include="')[1].split('"')[0]

    assert "[name='page']" in include
    assert 'name="page"' in page
