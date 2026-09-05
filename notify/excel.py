# -*- coding: utf-8 -*-
"""Gunluk ozetin eki — sezonun TAM listesi.

Mail yalnizca yonetici ozetini tasir; liste buraya gider. Ek KUMULATIFTIR:
"son 24 saatte ne oldu" degil, sezonun tamami. Sebep: alici dosyayi elle tutulan
"FA26 TRACKING LIST.xlsx"in yerine koyabilsin — o liste de sezonun tamamidir.

Uc sayfa — elle tutulan listedeki sekmelerin karsiligi:
  * "Tum Liste"        — her parca bir satir. Teslim edilenler YESIL boyanir.
  * "Teslim Edilenler" — yalnizca teslim edilmis parcalar.
  * "Sorunlular"       — GLS'in sorun bildirdikleri + hareketsizler + POD'u
                         eksikler ("Problematic Packages" sekmesi).
Daha fazla sayfa YOK: "hangi sekmedeydi" derdi baslar ve tek bir Excel filtresi
tum tabloyu tarayamaz. Ilk sayfa her seyi icerir, digerleri kolaylik.

`pandas` yerine dogrudan `openpyxl`: DataFrame'e ihtiyac yok, sutun genisligi,
dolgu ve baslik bicimi uzerinde denetim var. Ikisi de zaten kurulu.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import i18n
from notify import DOC_LANG
from tracking.db import explain_text, status_label

# Basliklar elle tutulan "FA26 TRACKING LIST.xlsx" ile BIREBIR ayni (buyuk harf,
# Ingilizce) — bu dosya onun yerini almali, alici ayni sutun adlarini gormeli.
# FLAG / STATUS / POD / LAST CHECK bizim eklerimiz. Elle listedeki
# "INVOICE NUMBER-2" ve "BRAND" sutunlari veritabaninda tutulmuyor, bu yuzden yok.
COLUMNS: tuple[tuple[str, str], ...] = (
    ("FLAG", "_flag"),
    ("PARCEL NUMBER", "tracking_no"),
    ("SHIPMENT AGENCY", "_agency"),
    ("SALES REP", "sales_rep"),
    ("STORE CODE", "store_code"),
    ("COUNTRY", "country"),
    ("INVOICE NUMBER", "invoice_number"),
    ("EMC INVOICES", "emc_invoices"),
    ("SHIPMENT DATE FROM TURKEY", "shipment_date"),
    ("STATUS", "_status_label"),
    ("EXPLAIN", "_explain"),
    ("DELIVERED DATE", "delivered_date"),
    ("SHIPMENT METHOD", "shipment_method"),
    ("POD", "_pod"),
    ("LAST CHECK", "last_checked_at"),
)

AGENCY = {"NL": "Gls Netherlands", "IE": "Gls Ireland"}

_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
# Elle tutulan listede teslim edilen satirlar yesile boyaniyor — ayni aliskanlik.
_DELIVERED_FILL = PatternFill("solid", fgColor="D9F2E4")


def _flags(data: dict) -> dict[str, str]:
    """Takip no -> "elle mudahale" etiketi. Tek etiket: onceligi yuksek olan.

    Bir parca hem sorunlu hem POD'suz olabilir; iki satir yazmak yerine en acil
    olani gosteriyoruz, digeri zaten STATUS/POD sutunlarindan okunur.
    """
    out: dict[str, str] = {}
    for label, key in (
        (i18n.t("excel.flag.delivered"), "delivered"),
        (i18n.t("excel.flag.not_handed_over"), "not_handed_over"),
        (i18n.t("excel.flag.pod_missing"), "pod_missing"),
        (i18n.t("excel.flag.stale", days=data["stale_days"]), "stale"),
        (i18n.t("excel.flag.partial"), "partial"),
        (i18n.t("excel.flag.damaged"), "damaged"),
        (i18n.t("excel.flag.exception"), "exceptions"),
        (i18n.t("excel.flag.returned"), "returned"),
    ):
        for row in data[key]:
            out[row["tracking_no"]] = label
    return out


def _value(key: str, parcel: dict, flag: str):
    if key == "_flag":
        return flag
    if key == "_agency":
        return AGENCY.get(parcel.get("channel", ""), parcel.get("channel", ""))
    if key == "_status_label":
        return status_label(parcel.get("status", ""))
    if key == "_pod":
        return i18n.t("common.yes") if parcel.get("pod_path") else i18n.t("common.none")
    if key == "_explain":
        # Elle listedeki EXPLAIN sutununun karsiligi. Sorun olayi (varsa) ONCE
        # gelir: son olay gecmisin ortasindaki sorunu gizleyebiliyor. Sonra
        # GLS'in son olay metni, Excel'den gelen eski aciklama, sorun sebebi.
        # `explain_text` iki is yapar: teslim notunu cevirir (GLS'in metinleri
        # zaten Ingilizce ama "teslim alan" satirini biz Turkce yaziyoruz, bu
        # dosya disariya gidiyor — bkz. `tracking/db.py:delivered_note`) ve
        # sona son islem tarihini ekler.
        return explain_text(parcel.get("issue_text") or parcel.get("last_event_text")
                            or parcel.get("explain") or parcel.get("problem_reason") or "",
                            parcel.get("last_event_at"))
    return parcel.get(key, "")


def _sheet(ws, parcels: list[dict], flags: dict[str, str], paint: bool) -> None:
    ws.append([title for title, _ in COLUMNS])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for parcel in parcels:
        flag = flags.get(parcel.get("tracking_no", ""), "")
        ws.append([_value(key, parcel, flag) for _, key in COLUMNS])
        if paint and parcel.get("status") == "delivered":
            for cell in ws[ws.max_row]:
                cell.fill = _DELIVERED_FILL

    # Baslik satiri sabit + otomatik filtre: 4000 satirlik bir tabloda ikisi de sart.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for i, (title, _) in enumerate(COLUMNS, start=1):
        longest = max([len(title)] + [
            len(str(ws.cell(row=r, column=i).value or ""))
            for r in range(2, min(ws.max_row, 200) + 1)
        ])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 40)


# Sorunlular sayfasina GIRMEYEN tur: `not_handed_over` (etiketi basilmis ama GLS
# koliyi henuz almamis). Kullanicinin karari — bunlar sorun degil, kamyonu
# bekleyen normal kolilerdir ve sayfayi yuzlerce satirla dolduruyorlardi.
# Tam listede (ilk sayfa) FLAG sutunundan hala okunabiliyorlar.
PROBLEM_SHEET_KINDS = ("exceptions", "damaged", "partial", "pod_missing", "stale")


def _problem_rows(data: dict) -> list[dict]:
    """Elle mudahale bekleyen tum parcalar — tek listede, tekrarsiz.

    `problem_parcels()` ayni parcayi birden fazla `problem_kind` ile
    dondurebilir (hem hareketsiz hem POD'suz olabilir); ilk gorulen kalir,
    cunku asagidaki sira zaten onem sirasidir.
    """
    rows: dict[str, dict] = {}
    for key in PROBLEM_SHEET_KINDS:
        for row in data[key]:
            rows.setdefault(row["tracking_no"], row)
    return list(rows.values())


def build(data: dict) -> bytes:
    """`digest.gather()` ciktisindan uc sayfalik xlsx uretir (daima Ingilizce)."""
    wb = Workbook()
    flags = _flags(data)
    parcels = data["parcels"]

    with i18n.use(DOC_LANG):
        ws = wb.active
        ws.title = i18n.t("excel.sheet.all")
        _sheet(ws, parcels, flags, paint=True)

        delivered = [p for p in parcels if p.get("status") == "delivered"]
        _sheet(wb.create_sheet(i18n.t("excel.sheet.delivered")), delivered, flags,
               paint=False)

        # Ucuncu sayfa elle tutulan listedeki "Problematic Packages" sekmesinin
        # karsiligi: GLS'in sorun bildirdikleri + gunlerdir hareketsiz kalanlar +
        # teslim edilip POD'u gelmeyenler. Hangisi oldugunu FLAG sutunu soyler,
        # sebebi ise EXPLAIN sutununda (GLS'in son olay metni; "resepsiyon kapali"
        # gibi aciklamalar oradan gelir).
        _sheet(wb.create_sheet(i18n.t("excel.sheet.problems")), _problem_rows(data),
               flags, paint=False)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(data: dict) -> str:
    season = data.get("season") or "GLS"
    with i18n.use(DOC_LANG):
        stem = i18n.t("excel.filename")
    return f"{season}_{stem}_{data['generated_at'].strftime('%Y-%m-%d')}.xlsx"
