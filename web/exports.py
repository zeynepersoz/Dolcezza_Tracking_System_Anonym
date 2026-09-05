# -*- coding: utf-8 -*-
"""Ekrandaki tabloyu Excel veya PDF dosyasina cevirir.

TEK jenerik cevirici: koli icerigi de takip listesi de buradan gecer. Her ekran
icin ayri bir yazici olsaydi biri duzeltilip digeri unutulurdu — POD belgesinde
(`web/pod_document.py`) yasandi, oradaki font secimi burada YENIDEN YAZILMAZ,
ayni fonksiyon cagrilir.

Cikti kullanicinin gordugu tablodur: baslik satiri + hucreler + otomatik filtre + dondurulmus ust satir.
Koli icerigi coklu indirmesinde 2. sayfa (İçerik) ayni kolideki satirlari pastel renkle gruplar.
"""
from __future__ import annotations

import io
import re

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import timez
from web.pod_document import font_pair

# Excel sayfa adi 31 karakteri asamaz ve `[]:*?/\` karakterlerini kabul etmez.
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

HEADER_FILL = PatternFill("solid", fgColor="1E293B")   # Dark slate-800
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TOTAL_FILL = PatternFill("solid", fgColor="E2E8F0")    # Slate-200

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

PASTEL_PALETTE = [
    PatternFill("solid", fgColor="EFF6FF"),  # Soft Sky / Blue (slate/sky 50)
    PatternFill("solid", fgColor="F0FDF4"),  # Soft Emerald / Mint (emerald 50)
    PatternFill("solid", fgColor="FAF5FF"),  # Soft Purple / Lavender (purple 50)
    PatternFill("solid", fgColor="FFFBEB"),  # Soft Amber / Warm (amber 50)
    PatternFill("solid", fgColor="FDF2F8"),  # Soft Rose / Pink (pink 50)
    PatternFill("solid", fgColor="F0FDFA"),  # Soft Teal / Cyan (teal 50)
]

INK = (30, 41, 59)
MUTED = (100, 116, 139)


def _cell(value) -> str:
    return "" if value is None else str(value)


def _value(value):
    """Excel hucresinin degeri. Sayilar SAYI kalir, metne cevrilmez.

    Her sey `str()`ten gecerken adet sutunu Excel'de metin oluyordu: alta
    toplam alinamiyor, sagdan hizalanmiyordu. Kullanici tam bunu \"exelde de
    sikinti cikariyor\" diye bildirdi. `bool` disarida tutulur — `int`in alt
    turudur ve hucreye 1/0 diye duserdi.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    return value if isinstance(value, (int, float)) else str(value)


def sheet_name(title: str) -> str:
    return _BAD_SHEET_CHARS.sub(" ", title).strip()[:31] or "Sheet1"


def _fill_sheet(sheet, headers: list[str], rows: list[list], is_lines_sheet: bool = False) -> None:
    """Bir sayfayi yazar: kalin baslik, otomatik filtre, dondurulmus ust satir, pastel gruplama."""
    sheet.append(headers)
    sheet.row_dimensions[1].height = 26
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.border = THIN_BORDER

    current_group_key = None
    color_idx = -1

    for row_idx, row in enumerate(rows, start=2):
        row_values = [_value(v) for v in row]
        sheet.append(row_values)
        sheet.row_dimensions[row_idx].height = 20

        is_total = bool(row and str(row[0]).strip().upper() in ("TOPLAM", "TOTAL"))

        if is_lines_sheet and not is_total:
            # 2. sayfa (Icerik): 1. kolon tracking_no / koli koduna gore pastel renk gruplamasi
            key = str(row[0]).strip() if row else ""
            if key != current_group_key:
                current_group_key = key
                color_idx = (color_idx + 1) % len(PASTEL_PALETTE)
            row_fill = PASTEL_PALETTE[color_idx]
        elif is_total:
            row_fill = TOTAL_FILL
        else:
            row_fill = None

        for col_idx, cell in enumerate(sheet[row_idx], start=1):
            cell.border = THIN_BORDER
            if row_fill:
                cell.fill = row_fill
            if is_total:
                cell.font = Font(bold=True)
            cell.alignment = Alignment(
                vertical="center",
                horizontal="center" if isinstance(cell.value, (int, float)) else "left",
            )

    # Otomatik sutun genislikleri
    for index, header in enumerate(headers, start=1):
        widest = max([len(_cell(header))]
                     + [len(_cell(r[index - 1])) for r in rows if index <= len(r)])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 4, 12), 48)

    # Ust satiri dondur
    sheet.freeze_panes = "A2"

    # Ust satira otomatik filtre ekle
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        last_col = get_column_letter(sheet.max_column)
        sheet.auto_filter.ref = f"A1:{last_col}{max(sheet.max_row, 2)}"


def to_xlsx_multi(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    """Cok sayfali `.xlsx`: [(sayfa adi, basliklar, satirlar), ...].

    Musteri bazli koli icerigi uc sayfa istiyor (ozet / satir satir / matris).
    2. sayfa (İçerik) ayni kolideki satirlari otomatik pastel renklere boyar.
    """
    book = Workbook()
    for index, (title, headers, rows) in enumerate(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = sheet_name(title)
        # 2. sayfa (index == 1) icerik/lines sayfasidir
        is_lines_sheet = (index == 1 and len(sheets) >= 2) or (title.lower() in ("i̇çerik", "icerik", "lines"))
        _fill_sheet(sheet, headers, rows, is_lines_sheet=is_lines_sheet)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def to_xlsx(headers: list[str], rows: list[list], title: str) -> bytes:
    """Tabloyu tek sayfali bir `.xlsx` olarak dondurur."""
    return to_xlsx_multi([(title, headers, rows)])


class _Table(FPDF):
    """Yatay A4 tablo. `title`/`table` adlari FPDF'in kendi uyeleridir, kullanilmaz."""

    def __init__(self, heading: str, subtitle: str):
        super().__init__(orientation="L", unit="mm", format="A4")
        self.heading_text = heading
        self.subtitle_text = subtitle
        regular, bold = font_pair()
        self.add_font("body", "", regular)
        self.add_font("body", "B", bold)
        self.set_auto_page_break(True, margin=14)

    def header(self):
        self.set_font("body", "B", 13)
        self.set_text_color(*INK)
        self.cell(0, 7, self.heading_text, new_x="LMARGIN", new_y="NEXT")
        if self.subtitle_text:
            self.set_font("body", "", 8.5)
            self.set_text_color(*MUTED)
            self.cell(0, 5, self.subtitle_text, new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def footer(self):
        self.set_y(-12)
        self.set_font("body", "", 7.5)
        self.set_text_color(*MUTED)
        self.cell(0, 5, f"{self.page_no()}", align="R")


def _draw_table(doc: "_Table", headers: list[str], rows: list[list]) -> None:
    """Kolon genisligi ICERIKTEN hesaplanir: sabit genislik verilirse uzun bir
    magaza adi kisa bir beden kolonuyla ayni yeri kaplar ve tablo okunmaz olur."""
    usable = doc.w - doc.l_margin - doc.r_margin
    weights = [
        max([len(_cell(headers[i]))] + [len(_cell(r[i])) for r in rows if i < len(r)])
        for i in range(len(headers))
    ]
    total = sum(weights) or 1
    widths = [max(usable * w / total, 12.0) for w in weights]
    scale = usable / sum(widths)
    widths = [w * scale for w in widths]

    doc.set_font("body", "B", 8)
    doc.set_text_color(*INK)
    doc.set_fill_color(241, 245, 249)
    for width, header in zip(widths, headers):
        doc.cell(width, 7, _cell(header), border="B", fill=True)
    doc.ln()

    doc.set_font("body", "", 8)
    for row in rows:
        for index, width in enumerate(widths):
            value = _cell(row[index]) if index < len(row) else ""
            doc.cell(width, 6, value, border="B")
        doc.ln()


def to_pdf(headers: list[str], rows: list[list], heading: str,
           subtitle: str = "") -> bytes:
    """Tabloyu yatay A4 PDF olarak dondurur."""
    doc = _Table(heading, subtitle)
    doc.add_page()
    _draw_table(doc, headers, rows)
    return bytes(doc.output())


def to_pdf_multi(heading: str, subtitle: str,
                 sections: list[tuple[str, list[str], list[list]]]) -> bytes:
    """Coklu tablo, TEK PDF — `to_xlsx_multi`nin PDF karsiligi.

    Her `(bolum basligi, basliklar, satirlar)` uclusu KENDI sayfasinda cizilir
    (kullanici istegi, 2026-08-31: "tüm raporların export edilebilir olması").
    """
    doc = _Table(heading, subtitle)
    for section_title, section_headers, section_rows in sections:
        doc.add_page()
        if section_title:
            doc.set_font("body", "B", 10)
            doc.set_text_color(*INK)
            doc.cell(0, 6, section_title, new_x="LMARGIN", new_y="NEXT")
            doc.ln(1)
        _draw_table(doc, section_headers, section_rows)
    return bytes(doc.output())


def safe_filename(value: str) -> str:
    """Magaza kodu gibi bir degeri dosya adina uygun hale getirir.

    Deger veritabanindan geliyor ama `Content-Disposition` basligina yaziliyor:
    tirnak, egik cizgi ve satir sonu oraya GIRMEMELI.
    """
    return re.sub(r"[^A-Za-z0-9._-]+", "-", value).strip("-")[:40]


def stamped_name(prefix: str, extension: str) -> str:
    """`takip-listesi-20260819-1435.xlsx` — ayni klasore inen dosyalar ezismesin."""
    return f"{prefix}-{timez.now().strftime('%Y%m%d-%H%M')}.{extension}"
